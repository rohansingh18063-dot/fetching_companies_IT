

import subprocess
import sys

REQUIRED = {"flask": "flask", "requests": "requests", "bs4": "beautifulsoup4", "dns": "dnspython",
            "ddgs": "ddgs", "openpyxl": "openpyxl"}
for module, package in REQUIRED.items():
    try:
        __import__(module)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", package])

import csv
import html as htmllib
import io
import os
import re
import sqlite3
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, wait
from urllib.parse import unquote, urljoin, urlparse
from urllib.robotparser import RobotFileParser

import dns.resolver
import requests
import urllib3
from bs4 import BeautifulSoup
from flask import Flask, Response, jsonify, request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "it_companies.db")
ROBOTS_NAME = "ITCompanyFinder"
BROWSER_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/130.0 Safari/537.36")
OSM_UA = "ITCompanyFinder/2.0 (contact-research)"
TIMEOUT = (5, 9)          # connect, read
MAX_BYTES = 1_500_000     # stop reading a page after 1.5 MB
WORKERS = 24              # companies checked at the same time
PAGE_WORKERS = 6          # pages of one company fetched at the same time
MAX_SUBPAGES = 9

# ---------------------------------------------------------------- data

STATES = {
    "Andhra Pradesh": ["Visakhapatnam", "Vijayawada", "Guntur", "Tirupati"],
    "Assam": ["Guwahati"],
    "Bihar": ["Patna"],
    "Chandigarh": ["Chandigarh"],
    "Chhattisgarh": ["Raipur", "Bhilai"],
    "Delhi": ["New Delhi", "Delhi"],
    "Goa": ["Panaji", "Margao"],
    "Gujarat": ["Ahmedabad", "Gandhinagar", "Surat", "Vadodara", "Rajkot"],
    "Haryana": ["Gurugram", "Faridabad", "Panchkula"],
    "Himachal Pradesh": ["Shimla"],
    "Jharkhand": ["Ranchi", "Jamshedpur"],
    "Karnataka": ["Bengaluru", "Mysuru", "Mangaluru", "Hubballi"],
    "Kerala": ["Kochi", "Thiruvananthapuram", "Kozhikode"],
    "Madhya Pradesh": ["Indore", "Bhopal", "Jabalpur", "Gwalior"],
    "Maharashtra": ["Mumbai", "Pune", "Navi Mumbai", "Thane", "Nagpur", "Nashik", "Aurangabad"],
    "Odisha": ["Bhubaneswar", "Cuttack"],
    "Punjab": ["Mohali", "Ludhiana", "Amritsar", "Jalandhar"],
    "Rajasthan": ["Jaipur", "Jodhpur", "Udaipur", "Kota"],
    "Tamil Nadu": ["Chennai", "Coimbatore", "Madurai", "Tiruchirappalli"],
    "Telangana": ["Hyderabad", "Warangal"],
    "Uttar Pradesh": ["Noida", "Greater Noida", "Lucknow", "Ghaziabad", "Kanpur", "Varanasi"],
    "Uttarakhand": ["Dehradun"],
    "West Bengal": ["Kolkata", "Durgapur"],
}

# Directory sites spell some cities differently
CITY_ALIASES = {
    "Bengaluru": ["bangalore"], "Gurugram": ["gurgaon", "gurugram"], "Mysuru": ["mysore", "mysuru"],
    "Mangaluru": ["mangalore", "mangaluru"], "Hubballi": ["hubli", "hubballi"],
    "Thiruvananthapuram": ["thiruvananthapuram", "trivandrum"], "Kozhikode": ["kozhikode", "calicut"],
    "Tiruchirappalli": ["tiruchirappalli", "trichy"], "New Delhi": ["delhi", "new-delhi"],
    "Vadodara": ["vadodara", "baroda"], "Visakhapatnam": ["visakhapatnam", "vizag"],
    "Panaji": ["goa", "panaji"], "Margao": ["goa"], "Aurangabad": ["aurangabad", "chhatrapati-sambhajinagar"],
}

SEARCH_QUERIES = [
    "IT companies in {city}, {state}",
    "software development company in {city}, {state}",
    "IT services company in {city}, {state}",
    "software company in {city}, {state}",
]
WEB_QUERIES = [
    "IT companies in {city}",
    "top software development companies in {city} list",
    "software company {city} {state} careers",
    "IT services company {city} hr email",
    "web and mobile app development company {city}",
]
DIRECTORY_URLS = [
    "https://www.goodfirms.co/it-services/{slug}",
    "https://www.goodfirms.co/it-services/{slug}?page=2",
    "https://www.goodfirms.co/it-services/{slug}?page=3",
    "https://www.goodfirms.co/directory/city/top-software-development-companies/{slug}",
    "https://www.goodfirms.co/directory/city/top-software-development-companies/{slug}?page=2",
    "https://www.goodfirms.co/directory/city/top-software-development-companies/{slug}?page=3",
    "https://www.goodfirms.co/directory/city/top-website-development-companies/{slug}",
    "https://www.goodfirms.co/directory/city/top-website-development-companies/{slug}?page=2",
    "https://www.mobileappdaily.com/directory/software-development-companies/in/{slug}",
    "https://www.mobileappdaily.com/directory/mobile-app-development-companies/in/{slug}",
]

# Parts of an email address (split on . _ - + digits) that mean HR / hiring
HR_TOKENS = {"hr", "hrd", "hrm", "hrs", "hrteam", "hrdept", "hrdesk", "hrmanager", "hrexecutive", "hrhead",
             "hradmin", "hrops", "hrindia", "career", "careers", "job", "jobs", "recruit", "recruiter",
             "recruiters", "recruitment", "recruitments", "recruiting", "talent", "talents",
             "talentacquisition", "ta", "hiring", "hire", "hireme", "people", "peopleops", "resume",
             "resumes", "cv", "cvs", "joinus", "join", "humanresource", "humanresources", "placement",
             "placements", "staffing", "openings", "apply", "vacancy", "vacancies", "intern", "internship",
             "fresher", "freshers", "campus", "work", "workwithus"}
HR_PREFIXES = ("career", "recruit", "talent", "hiring", "resume", "humanres", "joinus", "placement",
               "jobs", "hrteam", "hrdept", "hrmanager", "hrexec")
# Words right next to an address that show it is meant for job applications. Kept strict on purpose:
# footers say "Careers  Contact  info@x.com", so plain "career" nearby proves nothing.
HR_CONTEXT_RE = re.compile(r"resume|r[ée]sum[ée]|\bcvs?\b|\bhr\b|human resource|recruit|hiring|apply|"
                           r"send your (?:profile|application|details)|internship|job opening", re.I)
NOT_HR_LOCAL = ("sales", "support", "business", "marketing", "billing", "accounts", "partner", "press")
GENERAL_WORDS = ["info", "contact", "hello", "enquiry", "enquiries", "inquiry", "office", "admin", "mail",
                 "support", "sales", "business", "connect", "reach", "team", "biz", "query"]
CAREER_LINK_WORDS = ["career", "job", "join-us", "joinus", "join us", "join our", "hiring", "opening",
                     "vacanc", "work-with-us", "work with us", "life-at", "life at", "current-opening", "recruit"]
CONTACT_LINK_WORDS = ["contact", "reach-us", "reach us", "get-in-touch", "get in touch", "enquiry", "inquiry"]
CAREER_PATHS = ["/careers", "/career", "/jobs", "/join-us", "/current-openings"]
CONTACT_PATHS = ["/contact-us", "/contact"]
IT_WORDS = ["software", "it services", "information technology", "web development", "app development",
            "mobile app", "cloud", "saas", "devops", "erp", "crm", "digital transformation",
            "data analytics", "ai ", "artificial intelligence", "machine learning", "cyber security",
            "cybersecurity", "it solutions", "it consulting", "technology solutions", "website development",
            "custom software", "outsourcing", "offshore", "qa testing", "ui/ux", "blockchain", "developers",
            "web design", "ecommerce", "e-commerce", "android", "ios", "react", "python", "java", "php",
            "salesforce", "aws", "azure", "product engineering", "staff augmentation", "it company"]
TECH_NAME_RE = re.compile(r"soft|tech|info|digital|solution|system|labs?\b|web|app|cloud|data|cyber|dev|"
                          r"code|logic|byte|net|it\b|ai\b|consult", re.I)
FREE_MAIL = {"gmail.com", "yahoo.com", "yahoo.co.in", "hotmail.com", "outlook.com", "rediffmail.com",
             "live.com", "icloud.com", "aol.com", "protonmail.com", "ymail.com", "yahoo.in", "outlook.in"}
BAD_EMAIL_PARTS = ["example.", "domain.", "yourname", "youremail", "your-email", "email.com", "sentry",
                   "wixpress", "@2x", ".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".css", ".js",
                   "noreply", "no-reply", "donotreply", "test@", "user@", "name@", "username@", "abc@",
                   "xyz@", "@sentry", "godaddy", "schema.org", "@company.", "@yourdomain", "@website.",
                   "@mysite.", "@sitename.", "johndoe", "john.doe", "@email.", "placeholder"]
SKIP_DOMAINS = ["facebook.com", "linkedin.com", "instagram.com", "twitter.com", "x.com", "youtube.com",
                "justdial.com", "indiamart.com", "sulekha.com", "google.", "wikipedia.org", "wikimedia",
                "glassdoor", "naukri.com", "ambitionbox.com", "clutch.co", "goodfirms.co", "mobileappdaily",
                "techbehemoths", "designrush", "f6s.com", "sortlist", "g2.com", "capterra", "trustpilot",
                "crunchbase", "github.com", "medium.com", "pinterest", "apple.com", "microsoft.com",
                "amazon.", "wordpress.org", "wordpress.com", "w3.org", "cloudflare", "gstatic", "googleapis",
                "bit.ly", "t.me", "wa.me", "whatsapp", "telegram", "quora.com", "reddit.com", "tracxn",
                "zaubacorp", "tofler", "indiatimes", "timesofindia", "hindustantimes", "ndtv", "thehindu",
                "yourstory", "inc42", "forbes", "bloomberg", "blogspot", "behance", "dribbble", "upwork",
                "fiverr", "freelancer.", "indeed.", "monster", "shine.com", "foundit", "internshala",
                "instahyre", "cutshort", "wellfound", "angel.co", "yelp", "tripadvisor", "bing.com", "yahoo.",
                "duckduckgo", "gov.in", "nic.in", "ac.in", "edu.in", "openstreetmap", "mapquest", "hubspot",
                "gravatar", "jsdelivr", "unpkg", "typeform", "calendly", "tawk.to", "addtoany", "sharethis",
                "issuu", "slideshare", "scribd", "vimeo", "spotify", "play.google", "goo.gl", "tinyurl",
                "zoominfo", "rocketreach", "apollo.io", "lusha", "dnb.com", "manta.com", "bizdb", "grotal",
                "asklaila", "yellowpages", "europages", "kompass", "mouthshut", "tradeindia", "exportersindia"]
INDIA_RE = re.compile(r"\+91|\bindia\b|\bpvt\.? ?ltd\b|private limited", re.I)
LISTICLE_RE = re.compile(r"top|best|list|companies|directory|firms|agencies", re.I)

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.(?:[a-z]{2,24}|[A-Z]{2,24})")
OBFUSCATED_RE = re.compile(
    r"([A-Za-z0-9._%+\-]+)\s*(?:\[at\]|\(at\)|\{at\}|\[@\]|\(@\)|\s+at\s+(?=[a-z0-9\-]+\s*(?:\[dot\]|\(dot\))))\s*"
    r"([A-Za-z0-9\-]+(?:\s*(?:\[dot\]|\(dot\)|\{dot\}|\.)\s*[A-Za-z0-9\-]+)+)", re.I)
DOT_RE = re.compile(r"\s*(?:\[dot\]|\(dot\)|\{dot\})\s*", re.I)

# ---------------------------------------------------------------- database

db_lock = threading.Lock()
COLUMNS = {
    "domain": "TEXT PRIMARY KEY", "company": "TEXT", "website": "TEXT", "career_page": "TEXT",
    "contact_page": "TEXT", "email_1": "TEXT", "email_2": "TEXT", "email_3": "TEXT", "email_type": "TEXT",
    "email_source": "TEXT", "guessed_email": "TEXT", "phone": "TEXT", "address": "TEXT", "city": "TEXT",
    "state": "TEXT", "linkedin_page": "TEXT", "trust_score": "INTEGER", "notes": "TEXT", "source": "TEXT",
    "updated_at": "TEXT",
}


def db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False, timeout=30)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as conn:
        conn.execute("CREATE TABLE IF NOT EXISTS companies (%s)" % ", ".join(f"{k} {v}" for k, v in COLUMNS.items()))
        have = {r["name"] for r in conn.execute("PRAGMA table_info(companies)")}
        for k, v in COLUMNS.items():  # upgrade databases made by the older version
            if k not in have:
                conn.execute(f"ALTER TABLE companies ADD COLUMN {k} {v}")
        conn.execute("CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT)")
        conn.execute("""
            CREATE TABLE IF NOT EXISTS runs (
                id TEXT PRIMARY KEY, state TEXT, city TEXT, source TEXT,
                found INTEGER, saved INTEGER, skipped INTEGER, started_at TEXT
            )""")
        if "hr" not in {r["name"] for r in conn.execute("PRAGMA table_info(runs)")}:
            conn.execute("ALTER TABLE runs ADD COLUMN hr INTEGER DEFAULT 0")


def get_setting(key, default=""):
    with db() as conn:
        row = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return row["value"] if row else default


def set_setting(key, value):
    with db_lock, db() as conn:
        conn.execute("INSERT INTO settings (key, value) VALUES (?, ?) "
                     "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))


def save_row(row):
    cols = list(COLUMNS)
    with db_lock, db() as conn:
        conn.execute(
            f"INSERT INTO companies ({', '.join(cols)}) VALUES ({', '.join(':' + c for c in cols)}) "
            f"ON CONFLICT(domain) DO UPDATE SET {', '.join(f'{c}=excluded.{c}' for c in cols[1:])}",
            {c: row.get(c, "") for c in cols})


def existing_domains():
    with db() as conn:
        return {r["domain"] for r in conn.execute("SELECT domain FROM companies")}

# ---------------------------------------------------------------- helpers

session = requests.Session()
session.headers.update({"User-Agent": BROWSER_UA, "Accept-Language": "en-IN,en;q=0.9",
                        "Accept": "text/html,application/xhtml+xml,*/*;q=0.8"})
adapter = requests.adapters.HTTPAdapter(pool_connections=200, pool_maxsize=200, max_retries=0)
session.mount("http://", adapter)
session.mount("https://", adapter)


def root_domain(url):
    host = urlparse(url if "://" in url else "http://" + url).netloc.lower()
    host = host.split("@")[-1].split(":")[0]
    return host[4:] if host.startswith("www.") else host


def same_site(url, domain):
    d = root_domain(url)
    return d == domain or d.endswith("." + domain)


def skip_site(url):
    d = root_domain(url)
    return not d or "." not in d or any(s in d for s in SKIP_DOMAINS)


robots_cache = {}


def allowed(url):
    base = "{0.scheme}://{0.netloc}".format(urlparse(url))
    if base not in robots_cache:
        rp = RobotFileParser()
        try:
            r = session.get(base + "/robots.txt", timeout=(4, 5), verify=False)
            rp.parse(r.text.splitlines() if r.status_code == 200 else [])
        except requests.RequestException:
            rp.parse([])
        robots_cache[base] = rp
    return robots_cache[base].can_fetch(ROBOTS_NAME, url)


def fetch(url):
    """Return (final_url, html) or (None, None). Retries without SSL checks for broken certificates."""
    if not allowed(url):
        return None, None
    for verify in (True, False):
        try:
            with session.get(url, timeout=TIMEOUT, allow_redirects=True, verify=verify, stream=True) as r:
                ctype = r.headers.get("Content-Type", "text/html").lower()
                if r.status_code != 200 or "html" not in ctype:
                    return None, None
                body = b""
                for chunk in r.iter_content(65536):
                    body += chunk
                    if len(body) > MAX_BYTES:
                        break
                enc = r.encoding if "charset" in ctype else "utf-8"
                return r.url, body.decode(enc or "utf-8", errors="replace")
        except requests.exceptions.SSLError:
            continue
        except (requests.RequestException, UnicodeError, LookupError):
            return None, None
    return None, None


def decode_cfemail(hexstr):
    try:
        key = int(hexstr[:2], 16)
        return "".join(chr(int(hexstr[i:i + 2], 16) ^ key) for i in range(2, len(hexstr), 2))
    except ValueError:
        return ""


def clean_email(e):
    e = unquote(e).strip().strip(".,;:'\"()<>[]").lower()
    if not EMAIL_RE.fullmatch(e) or any(b in e for b in BAD_EMAIL_PARTS):
        return ""
    local = e.split("@")[0]
    if len(local) > 40 or re.fullmatch(r"[0-9a-f]{16,}", local):  # tracking ids, hashes
        return ""
    return e


def near_text(el):
    """Short text around an element: its own block if small, otherwise just the element."""
    parent = el.find_parent(["p", "li", "div", "td", "span"]) or el
    t = parent.get_text(" ")
    return t if len(t) < 300 else el.get_text(" ")


def extract_emails(html):
    """Return ({email: hr_context_bool}, soup). hr_context means the email sits next to hiring words."""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        if tag.name != "script" or tag.get("type") != "application/ld+json":
            tag.decompose()
    text = soup.get_text(" ")
    raw = {}

    def add(e, ctx=""):
        e = clean_email(e)
        if e:
            raw[e] = raw.get(e, False) or bool(HR_CONTEXT_RE.search(ctx))

    for a in soup.select("a[href^=mailto], a[href^=MAILTO]"):
        add(a["href"][7:].split("?")[0], near_text(a))
    for el in soup.select("[data-cfemail]"):
        add(decode_cfemail(el["data-cfemail"]), near_text(el))
    for m in EMAIL_RE.finditer(text):
        add(m.group(0), text[max(0, m.start() - 120):m.end() + 12])
    for m in EMAIL_RE.finditer(htmllib.unescape(html)):
        add(m.group(0))
    for m in OBFUSCATED_RE.finditer(text):
        add(f"{m.group(1)}@{DOT_RE.sub('.', m.group(2)).replace(' ', '')}",
            text[max(0, m.start() - 120):m.end() + 12])
    return raw, soup


mx_cache = {}


def has_mx(domain):
    if domain not in mx_cache:
        try:
            mx_cache[domain] = len(dns.resolver.resolve(domain, "MX", lifetime=5)) > 0
        except Exception:
            mx_cache[domain] = False
    return mx_cache[domain]


def is_hr_address(email):
    local = email.split("@")[0].lower()
    tokens = [t for t in re.split(r"[._\-+0-9]+", local) if t]
    if any(t in HR_TOKENS for t in tokens):
        return True
    return any(t.startswith(HR_PREFIXES) for t in tokens)


def classify(email, hr_context=False):
    if is_hr_address(email):
        return "HR"
    local = email.split("@")[0]
    if hr_context and not local.startswith(NOT_HR_LOCAL):
        return "HR"  # e.g. info@ printed under "send your resume to"
    if any(local.startswith(w) for w in GENERAL_WORDS):
        return "General"
    return "Other"


def rank_emails(emails, company_domain):
    """emails: {email: hr_context}. Return list of (email, kind) best first, only deliverable domains."""
    def score(item):
        e, ctx = item
        dom = e.split("@")[1]
        kind = classify(e, ctx)
        s = 100 if is_hr_address(e) else 80 if kind == "HR" else 40 if kind == "General" else 20
        if dom == company_domain or dom.endswith("." + company_domain) or company_domain.endswith("." + dom):
            s += 30
        if dom in FREE_MAIL:
            s -= 25
        return s
    good = [(e, c) for e, c in emails.items() if has_mx(e.split("@")[1])]
    good.sort(key=score, reverse=True)
    return [(e, classify(e, c)) for e, c in good]


def company_name(soup, fallback):
    for sel, attr in (("meta[property='og:site_name']", "content"), ("meta[name='application-name']", "content")):
        el = soup.select_one(sel)
        if el and el.get(attr, "").strip():
            return el[attr].strip()[:80]
    stem = re.sub(r"[^a-z0-9]", "", fallback.split(".")[0].lower())
    title = soup.title.get_text(" ").strip() if soup.title else ""
    parts = [p.strip() for p in re.split(r"\s*[|–—:]\s*|\s-\s", title) if p.strip()]
    for p in parts:  # the part of the title that matches the domain is the company name
        if stem[:5] and stem[:5] in re.sub(r"[^a-z0-9]", "", p.lower()):
            return p[:80]
    short = min(parts, key=len) if parts else ""
    generic = short.lower() in ("home", "homepage", "home page", "welcome", "index", "official website")
    return short[:80] if 0 < len(short) <= 30 and not generic else fallback.split(".")[0].title()

# ---------------------------------------------------------------- company discovery
# Every source calls add(candidate) as soon as it finds a website, so crawling starts immediately.


def google_places(city, state, api_key, add, stop, log):
    count = 0
    for q in SEARCH_QUERIES:
        token = None
        for _ in range(3):  # max 3 pages (60 results) per query
            if stop():
                return
            body = {"textQuery": q.format(city=city, state=state), "regionCode": "IN", "pageSize": 20}
            if token:
                body["pageToken"] = token
            r = requests.post(
                "https://places.googleapis.com/v1/places:searchText", json=body, timeout=20,
                headers={"X-Goog-Api-Key": api_key,
                         "X-Goog-FieldMask": "places.id,places.displayName,places.websiteUri,"
                                             "places.formattedAddress,places.nationalPhoneNumber,"
                                             "places.userRatingCount,places.businessStatus,nextPageToken"})
            if r.status_code != 200:
                log(f"Google Places error: {r.text[:200]}")
                return
            data = r.json()
            for p in data.get("places", []):
                if p.get("businessStatus") == "CLOSED_PERMANENTLY" or not p.get("websiteUri"):
                    continue
                count += 1
                add({"company": p.get("displayName", {}).get("text", ""), "website": p["websiteUri"],
                     "phone": p.get("nationalPhoneNumber", ""), "address": p.get("formattedAddress", ""),
                     "reviews": p.get("userRatingCount", 0), "source": "Google Places"})
            token = data.get("nextPageToken")
            if not token:
                break
            time.sleep(2)
    log(f"Google Places: {count} companies with websites")


def harvest_links(url, add, source, own_domain=None):
    """Open a directory / 'top companies' page and add every outbound company website on it."""
    final, html = fetch(url)
    if not html:
        return 0
    own = own_domain or root_domain(final)
    soup = BeautifulSoup(html, "html.parser")
    n = 0
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if not href.startswith("http") or same_site(href, own) or skip_site(href):
            continue
        name = a.get_text(" ").strip()
        if not name or len(name) > 60 or name.lower() in ("visit website", "website", "visit", "view website"):
            name = ""
        p = urlparse(href)
        n += add({"company": name, "website": f"{p.scheme}://{p.netloc}", "source": source})
    return n


def directories(city, add, stop, log):
    slugs = CITY_ALIASES.get(city, []) + [re.sub(r"[^a-z0-9]+", "-", city.lower()).strip("-")]
    slugs = list(dict.fromkeys(slugs))
    urls = [u.format(slug=s) for s in slugs for u in DIRECTORY_URLS]
    total = 0
    with ThreadPoolExecutor(6) as pool:
        for n in pool.map(lambda u: 0 if stop() else harvest_links(u, add, "Directory"), urls):
            total += n
    log(f"Business directories: {total} new company websites")


def web_search(city, state, add, stop, log):
    try:
        from ddgs import DDGS
    except ImportError:
        log("Web search unavailable (ddgs not installed)")
        return
    listicles, direct = [], 0
    for q in WEB_QUERIES:
        if stop():
            break
        try:
            results = DDGS(timeout=12).text(q.format(city=city, state=state), region="in-en", max_results=30)
        except Exception as exc:
            log(f"Web search: '{q.format(city=city, state=state)}' failed ({type(exc).__name__})")
            continue
        for res in results or []:
            url = res.get("href", "")
            if not url.startswith("http"):
                continue
            path = urlparse(url).path
            if skip_site(url) or LISTICLE_RE.search(path):
                listicles.append(url)       # a list page: harvest the companies it links to
            if not skip_site(url):
                direct += add({"company": res.get("title", "").split("|")[0].split(" - ")[0][:80],
                               "website": url, "source": "Web search"})
    listicles = [u for u in dict.fromkeys(listicles) if "goodfirms" not in u and "mobileappdaily" not in u]
    with ThreadPoolExecutor(6) as pool:
        harvested = sum(pool.map(lambda u: 0 if stop() else harvest_links(u, add, "Web search list"),
                                 listicles[:20]))
    log(f"Web search: {direct} direct + {harvested} from 'top companies' lists")


def openstreetmap(city, state, add, stop, log):
    h = {"User-Agent": OSM_UA}
    r = requests.get("https://nominatim.openstreetmap.org/search", headers=h, timeout=20,
                     params={"q": f"{city}, {state}, India", "format": "json", "limit": 1})
    hits = r.json()
    if not hits or stop():
        return
    s, n, w, e = hits[0]["boundingbox"]
    bbox = f"{s},{w},{n},{e}"
    query = f"""
    [out:json][timeout:40];
    (
      nwr["office"~"it|software|telecommunication|company"]["website"]({bbox});
      nwr["office"]["contact:website"]({bbox});
      nwr["company"~"software|it",i]({bbox});
      nwr["name"~"infotech|software|technologies|infosys|solutions",i]["website"]({bbox});
    );
    out center tags;"""
    r = requests.post("https://overpass-api.de/api/interpreter", data={"data": query}, headers=h, timeout=50)
    count = 0
    for el in r.json().get("elements", []):
        t = el.get("tags", {})
        site = t.get("website") or t.get("contact:website")
        if not site:
            continue
        if not site.startswith("http"):
            site = "http://" + site
        addr = ", ".join(x for x in [t.get("addr:housenumber"), t.get("addr:street"), t.get("addr:city")] if x)
        count += add({"company": t.get("name", ""), "website": site,
                      "phone": t.get("phone") or t.get("contact:phone", ""), "address": addr,
                      "source": "OpenStreetMap"})
    log(f"OpenStreetMap: {count} new company websites")


# ---------------------------------------------------------------- website crawl


def crawl_company(c, city, state, guess=True):
    domain = root_domain(c["website"])
    row = {"domain": domain, "company": c.get("company", ""), "website": c["website"], "career_page": "",
           "contact_page": "", "email_1": "", "email_2": "", "email_3": "", "email_type": "", "email_source": "",
           "guessed_email": "", "phone": c.get("phone", ""), "address": c.get("address", ""), "city": city,
           "state": state, "linkedin_page": "", "trust_score": 0, "notes": "", "source": c.get("source", ""),
           "updated_at": time.strftime("%Y-%m-%d %H:%M"), "is_it": False}

    home_url, html = fetch(f"https://{domain}")
    if not html:
        home_url, html = fetch(f"http://{domain}")
    if not html and urlparse(c["website"]).netloc.lower() != domain:
        home_url, html = fetch(c["website"])
    if not html:
        row["notes"] = "Website not reachable"
        return row
    domain = root_domain(home_url) or domain   # follow redirects to the real domain
    row["domain"], row["website"] = domain, home_url

    emails, soup = extract_emails(html)
    home_emails = set(emails)            # footer addresses repeat on every page
    source_of = {e: home_url for e in emails}
    meta = " ".join(m.get("content", "") for m in soup.select("meta[name=description], meta[name=keywords]"))
    page_text = (soup.get_text(" ") + " " + meta).lower()
    it_hits = sum(1 for w in IT_WORDS if w in page_text)
    row["company"] = row["company"] or company_name(soup, domain)
    tech_name = bool(TECH_NAME_RE.search(row["company"] + " " + domain))
    row["is_it"] = it_hits >= 3 or (it_hits >= 1 and tech_name)
    city_hit = city.lower() in page_text or any(a.replace("-", " ") in page_text for a in CITY_ALIASES.get(city, []))
    india_hit = domain.endswith(".in") or bool(INDIA_RE.search(page_text))

    careers, contacts, others = [], [], []
    for a in soup.find_all("a", href=True):
        href = urljoin(home_url, a["href"]).split("#")[0]
        if href.startswith("tel:") and not row["phone"]:
            row["phone"] = unquote(href[4:]).strip()
            continue
        if "linkedin.com/company" in href and not row["linkedin_page"]:
            row["linkedin_page"] = href
        if not href.startswith("http") or not same_site(href, domain):
            continue
        label = (a.get_text(" ") + " " + href).lower()
        if any(w in label for w in CAREER_LINK_WORDS):
            careers.append(href)
        elif any(w in label for w in CONTACT_LINK_WORDS):
            contacts.append(href)
        elif "about" in label or "team" in label:
            others.append(href)
    careers = list(dict.fromkeys(careers))
    contacts = list(dict.fromkeys(contacts))
    if not careers:
        careers = [urljoin(home_url, p) for p in CAREER_PATHS]
    if not contacts:
        contacts = [urljoin(home_url, p) for p in CONTACT_PATHS]
    queue = careers[:4] + contacts[:3] + others[:2]
    queue = [u for u in dict.fromkeys(queue) if u.rstrip("/") != home_url.rstrip("/")][:MAX_SUBPAGES]
    career_set = set(careers)

    def visit(link):
        final, sub = fetch(link)
        if not sub:
            return link, None, {}, None
        is_career = link in career_set or any(w in final.lower() for w in ("career", "job", "join", "hiring", "opening"))
        found, sub_soup = extract_emails(sub)
        if is_career:  # an address printed only on the careers page (not the site-wide footer) is the HR inbox
            for e in found:
                if e not in home_emails and not e.split("@")[0].startswith(NOT_HR_LOCAL):
                    found[e] = True
        return link, final, found, (sub_soup if is_career else None), sub_soup.get_text(" ").lower()

    def merge(results):
        nonlocal city_hit, india_hit
        more = []
        for res in results:
            link, final, found, career_soup = res[:4]
            if not final:
                continue
            text = res[4]
            city_hit = city_hit or city.lower() in text
            india_hit = india_hit or bool(INDIA_RE.search(text))
            for e, ctx in found.items():
                if ctx and not emails.get(e):
                    source_of[e] = final       # point to the page that says "send your resume to ..."
                emails[e] = emails.get(e, False) or ctx
                source_of.setdefault(e, final)
            if career_soup is not None:
                if not row["career_page"]:
                    row["career_page"] = final
                for a in career_soup.find_all("a", href=True):   # second hop: individual job / apply pages
                    href = urljoin(final, a["href"]).split("#")[0]
                    lab = (a.get_text(" ") + " " + href).lower()
                    if same_site(href, domain) and any(w in lab for w in ("apply", "opening", "position", "job", "vacanc", "hr")):
                        more.append(href)
            elif link in contacts and not row["contact_page"]:
                row["contact_page"] = final
        return more

    with ThreadPoolExecutor(PAGE_WORKERS) as pool:
        more = merge(pool.map(visit, queue))
        if not any(is_hr_address(e) or c for e, c in emails.items()):
            seen = set(queue) | {home_url}
            more = [u for u in dict.fromkeys(more) if u not in seen][:4]
            merge(pool.map(visit, more))

    ranked = rank_emails(emails, domain)[:3]
    for i, (e, _) in enumerate(ranked, 1):
        row[f"email_{i}"] = e
    kinds = [k for _, k in ranked]
    row["email_type"] = "HR" if "HR" in kinds else ("General" if ranked else "")
    if ranked:
        best = next((e for e, k in ranked if k == "HR"), ranked[0][0])
        row["email_source"] = source_of.get(best, "")
    domain_mx = has_mx(domain)
    if guess and row["email_type"] != "HR" and domain_mx:
        row["guessed_email"] = f"hr@{domain}, careers@{domain}"
        if not ranked:
            row["email_type"] = "Guess"

    # Trust score: does this look like a real, active IT company in this city?
    score = 20                                  # website loads
    score += min(it_hits, 5) * 6                # IT keywords on homepage
    score += 10 if domain_mx else 0             # domain receives email
    score += 10 if any(same_site("http://" + e.split("@")[1], domain) for e, _ in ranked) else 0
    score += 5 if row["linkedin_page"] else 0
    score += 10 if row["career_page"] else 0
    score += 10 if city_hit else 0
    # Directory / Places / OSM listings are already city-specific; web search results can be from anywhere
    row["is_local"] = city_hit or india_hit or not row["source"].startswith("Web search")
    score += 5 if c.get("reviews", 0) >= 5 else 0
    row["trust_score"] = min(score, 100)

    notes = []
    if not row["is_it"]:
        notes.append("Few IT signals - check manually")
    if not city_hit:
        notes.append(f"{city} not mentioned on site")
    if row["email_type"] == "HR":
        best_ctx = next((e for e, k in ranked if k == "HR"), "")
        if best_ctx and not is_hr_address(best_ctx):
            notes.append("Listed for resumes/hiring on the site")
    elif ranked:
        notes.append("No HR email published - using general email")
    if row["guessed_email"]:
        notes.append("hr@/careers@ is a guess - verify before mass mailing")
    if any(e.split("@")[1] in FREE_MAIL for e, _ in ranked):
        notes.append("Uses free email provider")
    row["notes"] = "; ".join(notes)
    return row

# ---------------------------------------------------------------- jobs

jobs = {}


def run_job(job_id, state, city, api_key, limit, min_trust, skip_known, hr_only, guess):
    job = jobs[job_id]
    lock = threading.Lock()

    def log(msg):
        job["log"].append(time.strftime("%H:%M:%S ") + msg)

    stop = lambda: job["stop"] or job["total"] >= limit
    known = existing_domains() if skip_known else set()
    seen = set()
    pool = ThreadPoolExecutor(max_workers=WORKERS)
    futures = []

    def process(c):
        if job["stop"]:
            return
        try:
            row = crawl_company(c, city, state, guess)
        except Exception as exc:  # one bad site must not stop the run
            row = None
            log(f"Skipped {c['website']}: {type(exc).__name__}")
        with lock:
            job["done"] += 1
            if not row:
                job["skipped"] += 1
                return
            if row["domain"] in job["saved_domains"]:
                return
            keep = (row["is_it"] and row.get("is_local") and row["trust_score"] >= min_trust
                    and (row["email_type"] == "HR" if hr_only else bool(row["email_1"] or row["guessed_email"])))
            if not keep:
                job["skipped"] += 1
                reason = ("not reachable" if row["notes"] == "Website not reachable"
                          else "not an IT company" if not row["is_it"]
                          else "not in India" if not row.get("is_local")
                          else "no HR email" if hr_only and row["email_type"] != "HR"
                          else "no email" if not (row["email_1"] or row["guessed_email"]) else "low trust")
                job["skip_reasons"][reason] = job["skip_reasons"].get(reason, 0) + 1
                return
            row.pop("is_it", None)
            row.pop("is_local", None)
            job["saved_domains"].add(row["domain"])
            save_row(row)
            job["rows"].append(row)
            if row["email_type"] == "HR":
                job["hr"] += 1
        log(f"Saved  {row['company']} - {row['email_1'] or row['guessed_email'].split(',')[0] + ' (guess)'}"
            f"{'  [HR]' if row['email_type'] == 'HR' else ''}")

    def add(c):
        d = root_domain(c.get("website", ""))
        with lock:
            if not d or d in seen or skip_site(c["website"]) or stop():
                return 0
            seen.add(d)
            if d in known:
                job["known"] += 1
                return 0
            job["total"] += 1
            futures.append(pool.submit(process, c))
        return 1

    def safe(name, fn):
        try:
            fn()
        except Exception as exc:
            log(f"{name} failed: {type(exc).__name__}: {str(exc)[:120]}")

    sources = []
    if api_key:
        sources.append(("Google Places", lambda: google_places(city, state, api_key, add, stop, log)))
    sources += [("Directories", lambda: directories(city, add, stop, log)),
                ("Web search", lambda: web_search(city, state, add, stop, log)),
                ("OpenStreetMap", lambda: openstreetmap(city, state, add, stop, log))]
    source_names = " + ".join(n for n, _ in sources)
    log(f"Searching IT companies in {city}, {state} using {source_names} ...")
    try:
        with ThreadPoolExecutor(len(sources)) as dpool:
            list(dpool.map(lambda s: safe(*s), sources))
        job["discovery_done"] = True
        if job["known"]:
            log(f"Skipped {job['known']} companies already in your database")
        if not job["total"]:
            log("No new companies found. Try another city, untick 'Skip companies already in database', "
                "or add a Google Places key in Settings.")
        else:
            log(f"Discovery finished: {job['total']} company websites to check")
        while True:  # futures can still be added while discovery threads finish
            with lock:
                pending = list(futures)
            wait(pending)
            with lock:
                if len(futures) == len(pending):
                    break
        reasons = ", ".join(f"{v} {k}" for k, v in sorted(job["skip_reasons"].items(), key=lambda x: -x[1]))
        log(f"Finished. Saved {len(job['rows'])} companies ({job['hr']} with HR email)."
            + (f" Skipped: {reasons}." if reasons else ""))
    except Exception as exc:
        log(f"Error: {exc}")
    pool.shutdown(wait=False, cancel_futures=True)
    job["status"] = "done"
    with db_lock, db() as conn:
        conn.execute("INSERT OR REPLACE INTO runs (id, state, city, source, found, saved, skipped, started_at, hr) "
                     "VALUES (?,?,?,?,?,?,?,?,?)",
                     (job_id, state, city, source_names, job["total"], len(job["rows"]), job["skipped"],
                      job["started_at"], job["hr"]))

# ---------------------------------------------------------------- web app

app = Flask(__name__)
EXPORT_COLUMNS = ["company", "website", "email_type", "email_1", "email_2", "email_3", "guessed_email",
                  "email_source", "career_page", "contact_page", "phone", "address", "city", "state",
                  "linkedin_page", "trust_score", "notes", "source", "updated_at"]
EXPORT_HEADERS = ["Company", "Website", "Email Type", "Best Email (HR first)", "Email 2", "Email 3",
                  "Likely HR Email (guess)", "Email Found On", "Career Page", "Contact Page", "Phone",
                  "Address", "City", "State", "LinkedIn Page", "Trust Score", "Notes", "Source", "Updated"]


def sort_rows(rows):
    order = {"HR": 0, "General": 1, "Guess": 2}
    return sorted(rows, key=lambda r: (order.get(r.get("email_type"), 3), -(r.get("trust_score") or 0)))


def to_csv(rows, filename):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(EXPORT_HEADERS)
    for r in sort_rows(rows):
        w.writerow([r.get(c, "") for c in EXPORT_COLUMNS])
    return Response("﻿" + buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={filename}.csv"})


def to_xlsx(rows, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "IT companies"
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    green = PatternFill("solid", fgColor="DCFCE7")
    amber = PatternFill("solid", fgColor="FEF3C7")
    link_cols = {EXPORT_COLUMNS.index(c) for c in ("website", "career_page", "contact_page", "linkedin_page")}
    for r in sort_rows(rows):
        ws.append([r.get(c, "") for c in EXPORT_COLUMNS])
        line = ws[ws.max_row]
        for i in link_cols:
            if str(line[i].value or "").startswith("http"):
                line[i].hyperlink = line[i].value
                line[i].font = Font(color="2563EB", underline="single")
        fill = green if r.get("email_type") == "HR" else amber if r.get("email_type") == "Guess" else None
        if fill:
            line[2].fill = line[3].fill = fill
    widths = [28, 30, 11, 32, 28, 28, 34, 30, 30, 30, 16, 36, 14, 16, 30, 8, 46, 16, 16]
    for i, wdt in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wdt
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={filename}.xlsx"})


def export(rows, filename, fmt):
    return to_xlsx(rows, filename) if fmt == "xlsx" else to_csv(rows, filename)


def filtered_rows(args, limit=None, offset=0):
    sql, params = " FROM companies WHERE 1=1", []
    if args.get("q"):
        sql += (" AND (company LIKE ? OR website LIKE ? OR email_1 LIKE ? OR email_2 LIKE ? OR email_3 LIKE ?"
                " OR guessed_email LIKE ?)")
        params += [f"%{args['q']}%"] * 6
    for field in ("state", "city", "email_type"):
        if args.get(field):
            sql += f" AND {field} = ?"
            params.append(args[field])
    if args.get("min_trust"):
        sql += " AND trust_score >= ?"
        params.append(int(args["min_trust"]))
    with db() as conn:
        total = conn.execute("SELECT COUNT(*)" + sql, params).fetchone()[0]
        q = ("SELECT *" + sql + " ORDER BY CASE email_type WHEN 'HR' THEN 0 WHEN 'General' THEN 1 "
             "WHEN 'Guess' THEN 2 ELSE 3 END, updated_at DESC, trust_score DESC")
        if limit:
            q += f" LIMIT {int(limit)} OFFSET {int(offset)}"
        rows = [dict(r) for r in conn.execute(q, params)]
    return total, rows


def flag(d, key, default=False):
    v = d.get(key, default)
    return v if isinstance(v, bool) else str(v).lower() in ("1", "true", "yes", "on")


@app.post("/api/start")
def start():
    d = request.json
    if any(j["status"] == "running" for j in jobs.values()):
        return jsonify(error="A search is already running. Please wait for it to finish."), 409
    job_id = uuid.uuid4().hex[:8]
    jobs[job_id] = {"id": job_id, "status": "running", "log": [], "rows": [], "done": 0, "total": 0,
                    "skipped": 0, "known": 0, "hr": 0, "stop": False, "discovery_done": False,
                    "skip_reasons": {}, "saved_domains": set(), "city": d["city"].strip(), "state": d["state"],
                    "started_at": time.strftime("%Y-%m-%d %H:%M")}
    api_key = get_setting("google_key") or os.getenv("GOOGLE_PLACES_KEY", "")
    threading.Thread(target=run_job, daemon=True, args=(
        job_id, d["state"], d["city"].strip(), api_key, max(5, min(int(d.get("limit") or 100), 500)),
        int(d.get("min_trust") or 0), flag(d, "skip_known"), flag(d, "hr_only"), flag(d, "guess", True))).start()
    return jsonify(job_id=job_id)


def job_json(j):
    return {k: v for k, v in j.items() if k != "saved_domains"} | {"rows": sort_rows(j["rows"])}


@app.get("/api/job/active")
def active_job():
    for j in jobs.values():
        if j["status"] == "running":
            return jsonify(job_json(j))
    return jsonify(None)


@app.get("/api/job/<job_id>")
def job_status(job_id):
    j = jobs.get(job_id)
    return jsonify(job_json(j)) if j else (jsonify(error="not found"), 404)


@app.post("/api/job/<job_id>/stop")
def job_stop(job_id):
    if job_id in jobs:
        jobs[job_id]["stop"] = True
    return jsonify(ok=True)


@app.get("/api/job/<job_id>/export.<fmt>")
def job_export(job_id, fmt):
    j = jobs[job_id]
    rows = [r for r in j["rows"] if r["email_type"] == "HR"] if request.args.get("hr_only") else j["rows"]
    name = f"it_companies_{j['city'].replace(' ', '_').lower()}" + ("_hr_only" if request.args.get("hr_only") else "")
    return export(rows, name, fmt)


@app.get("/api/companies")
def companies():
    page = max(int(request.args.get("page", 1)), 1)
    total, rows = filtered_rows(request.args, limit=50, offset=(page - 1) * 50)
    return jsonify(total=total, rows=rows, page=page, pages=max((total + 49) // 50, 1))


@app.get("/api/companies.<fmt>")
def companies_export(fmt):
    _, rows = filtered_rows(request.args)
    return export(rows, "it_companies", fmt)


@app.get("/api/hr-emails")
def hr_emails():
    _, rows = filtered_rows(request.args)
    return jsonify(emails=[r["email_1"] for r in rows if r["email_type"] == "HR" and r["email_1"]])


@app.delete("/api/companies/<path:domain>")
def delete_company(domain):
    with db_lock, db() as conn:
        conn.execute("DELETE FROM companies WHERE domain = ?", (domain,))
    return jsonify(ok=True)


@app.get("/api/filters")
def filters():
    with db() as conn:
        pairs = conn.execute("SELECT DISTINCT state, city FROM companies ORDER BY state, city").fetchall()
    out = {}
    for p in pairs:
        out.setdefault(p["state"], []).append(p["city"])
    return jsonify(out)


@app.get("/api/stats")
def stats():
    with db() as conn:
        one = lambda q: conn.execute(q).fetchone()[0]
        data = {
            "total": one("SELECT COUNT(*) FROM companies"),
            "hr": one("SELECT COUNT(*) FROM companies WHERE email_type='HR'"),
            "career": one("SELECT COUNT(*) FROM companies WHERE career_page != ''"),
            "cities": one("SELECT COUNT(DISTINCT city) FROM companies"),
            "top_cities": [dict(r) for r in conn.execute(
                "SELECT city, state, COUNT(*) n, SUM(email_type='HR') hr FROM companies "
                "GROUP BY city, state ORDER BY n DESC LIMIT 8")],
            "runs": [dict(r) for r in conn.execute("SELECT * FROM runs ORDER BY started_at DESC LIMIT 50")],
        }
    return jsonify(data)


@app.get("/api/settings")
def read_settings():
    key = get_setting("google_key")
    return jsonify(google_key_set=bool(key), google_key_hint=("..." + key[-4:]) if key else "",
                   limit=int(get_setting("limit", "100")), min_trust=int(get_setting("min_trust", "40")))


@app.post("/api/settings")
def write_settings():
    d = request.json
    if "google_key" in d:
        set_setting("google_key", d["google_key"].strip())
    for k in ("limit", "min_trust"):
        if k in d:
            set_setting(k, str(int(d[k])))
    return read_settings()


@app.get("/api/states")
def states():
    return jsonify(STATES)


@app.get("/")
def index():
    return PAGE


PAGE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>HR Finder</title>
<style>
:root{--bg:#f4f5f7;--card:#fff;--ink:#111827;--muted:#6b7280;--line:#e5e7eb;--accent:#2563eb;--accent-soft:#eff4ff;
--ok:#15803d;--ok-soft:#ecfdf3;--warn:#b45309;--warn-soft:#fff7ed;--side:#0f172a;--side-ink:#cbd5e1}
*{box-sizing:border-box}html,body{height:100%}
body{margin:0;font:14px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;background:var(--bg);color:var(--ink);display:flex}
aside{width:220px;flex:none;background:var(--side);color:var(--side-ink);padding:20px 12px;display:flex;flex-direction:column;gap:4px}
.brand{color:#fff;font-weight:700;font-size:17px;padding:4px 10px 18px}.brand small{display:block;font-weight:400;font-size:12px;color:#94a3b8}
nav a{display:flex;gap:10px;align-items:center;padding:9px 10px;border-radius:7px;color:var(--side-ink);text-decoration:none;cursor:pointer}
nav a:hover{background:#1e293b}nav a.on{background:var(--accent);color:#fff}
nav svg{width:17px;height:17px;flex:none}
.running-pill{margin-top:auto;font-size:12px;background:#1e293b;border-radius:7px;padding:10px;display:none;cursor:pointer}
.running-pill b{color:#fff}
main{flex:1;min-width:0;overflow:auto}
.wrap{max-width:1280px;margin:0 auto;padding:28px 24px}
h1{font-size:21px;margin:0 0 2px}.sub{color:var(--muted);margin:0 0 20px}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:18px;margin-bottom:16px}
.card h2{font-size:15px;margin:0 0 12px}
.tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin-bottom:16px}
.tile{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:14px 16px}
.tile b{font-size:26px;display:block;line-height:1.2}.tile span{color:var(--muted);font-size:13px}
.tile.hr b{color:var(--ok)}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}
.two{display:grid;grid-template-columns:1fr 1fr;gap:16px}
label{display:block;font-size:12px;color:var(--muted);margin-bottom:4px;font-weight:500}
input,select{width:100%;padding:9px 10px;border:1px solid var(--line);border-radius:7px;font:inherit;background:#fff;color:var(--ink)}
input:focus,select:focus{outline:2px solid var(--accent-soft);border-color:var(--accent)}
.btn{padding:9px 15px;border:1px solid transparent;border-radius:7px;background:var(--accent);color:#fff;font:inherit;font-weight:600;cursor:pointer;text-decoration:none;display:inline-flex;gap:6px;align-items:center}
.btn.sec{background:#fff;color:var(--ink);border-color:var(--line)}.btn.sm{padding:5px 9px;font-size:12px}
.btn.ok{background:var(--ok)}
.btn.danger{background:#fff;color:#b91c1c;border-color:#fecaca}.btn:disabled{opacity:.5;cursor:default}
.row{display:flex;gap:10px;flex-wrap:wrap;align-items:center}.mt{margin-top:14px}.right{margin-left:auto}
.check{display:flex;gap:6px;align-items:center;color:var(--ink);font-size:14px;margin:0;font-weight:400}.check input{width:auto}
.opts{display:flex;gap:8px 20px;flex-wrap:wrap;margin-top:14px}
.bar{height:8px;background:var(--line);border-radius:4px;overflow:hidden;margin:12px 0}.bar i{display:block;height:100%;background:var(--accent);width:0;transition:width .4s}
.counts{display:flex;gap:18px;flex-wrap:wrap;font-size:13px;color:var(--muted)}.counts b{color:var(--ink);font-size:15px}.counts .g b{color:var(--ok)}
.log{font:12px/1.55 ui-monospace,Menlo,Consolas,monospace;background:#0f172a;color:#cbd5e1;border-radius:8px;padding:10px 12px;height:150px;overflow:auto;white-space:pre-wrap;margin-top:12px}
details summary{cursor:pointer;color:var(--muted);font-size:13px;margin-top:10px}
.tbl{overflow:auto;border:1px solid var(--line);border-radius:8px;max-height:70vh}
table{border-collapse:collapse;width:100%;font-size:13px;background:#fff}
th,td{text-align:left;padding:8px 10px;border-bottom:1px solid var(--line);vertical-align:top}
th{color:var(--muted);font-weight:600;background:#fafafa;position:sticky;top:0;white-space:nowrap;z-index:1}
tr:last-child td{border-bottom:0}a{color:var(--accent)}
tr.is-hr td:first-child{box-shadow:inset 3px 0 0 var(--ok)}
.badge{display:inline-block;padding:0 7px;border-radius:99px;font-size:11px;font-weight:700;line-height:18px}
.b-hr{background:var(--ok-soft);color:var(--ok)}.b-gen{background:#f3f4f6;color:#4b5563}.b-guess{background:var(--warn-soft);color:var(--warn)}
.score{font-weight:600}.s-hi{color:var(--ok)}.s-mid{color:var(--warn)}.s-lo{color:#b91c1c}
.email{display:flex;gap:6px;align-items:center;white-space:nowrap;margin-bottom:3px}
.email.hr b{color:var(--ok)}
.guess{color:var(--warn);font-style:italic;white-space:nowrap}
.copy{border:0;background:none;color:var(--muted);cursor:pointer;padding:0;font-size:12px}.copy:hover{color:var(--accent)}
.muted{color:var(--muted)}.note{font-size:12px;color:var(--muted);max-width:240px}.small{font-size:12px}
.empty{text-align:center;color:var(--muted);padding:30px}
.pager{display:flex;gap:8px;align-items:center;justify-content:flex-end;margin-top:12px}
.toast{position:fixed;bottom:20px;right:20px;background:#111827;color:#fff;padding:10px 14px;border-radius:8px;opacity:0;transition:opacity .2s;pointer-events:none;z-index:9}
.toast.on{opacity:1}
.hint{font-size:12px;color:var(--muted);margin:8px 0 0}
.dlbar{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin:0 0 10px}
section{display:none}section.on{display:block}
@media(max-width:800px){body{flex-direction:column}aside{width:auto;flex-direction:row;overflow:auto;padding:10px}.brand,.running-pill{display:none!important}nav{display:flex;gap:4px}.two{grid-template-columns:1fr}.wrap{padding:18px 14px}}
</style></head><body>
<aside>
  <div class="brand">HR Finder<small>IT company database</small></div>
  <nav>
    <a data-p="dash"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><rect x="3" y="3" width="7" height="9"/><rect x="14" y="3" width="7" height="5"/><rect x="14" y="12" width="7" height="9"/><rect x="3" y="16" width="7" height="5"/></svg>Dashboard</a>
    <a data-p="search"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><circle cx="11" cy="11" r="7"/><path d="m20 20-3.5-3.5"/></svg>New Search</a>
    <a data-p="db"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><ellipse cx="12" cy="5" rx="8" ry="3"/><path d="M4 5v14c0 1.7 3.6 3 8 3s8-1.3 8-3V5"/><path d="M4 12c0 1.7 3.6 3 8 3s8-1.3 8-3"/></svg>Database</a>
    <a data-p="history"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 2"/></svg>History</a>
    <a data-p="settings"><svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><circle cx="12" cy="12" r="3"/><path d="M19.4 15a1.7 1.7 0 0 0 .3 1.8l.1.1a2 2 0 1 1-2.8 2.8l-.1-.1a1.7 1.7 0 0 0-1.8-.3 1.7 1.7 0 0 0-1 1.5V21a2 2 0 1 1-4 0v-.1a1.7 1.7 0 0 0-1.1-1.5 1.7 1.7 0 0 0-1.8.3l-.1.1a2 2 0 1 1-2.8-2.8l.1-.1a1.7 1.7 0 0 0 .3-1.8 1.7 1.7 0 0 0-1.5-1H3a2 2 0 1 1 0-4h.1a1.7 1.7 0 0 0 1.5-1.1 1.7 1.7 0 0 0-.3-1.8l-.1-.1a2 2 0 1 1 2.8-2.8l.1.1a1.7 1.7 0 0 0 1.8.3H9a1.7 1.7 0 0 0 1-1.5V3a2 2 0 1 1 4 0v.1a1.7 1.7 0 0 0 1 1.5 1.7 1.7 0 0 0 1.8-.3l.1-.1a2 2 0 1 1 2.8 2.8l-.1.1a1.7 1.7 0 0 0-.3 1.8V9a1.7 1.7 0 0 0 1.5 1H21a2 2 0 1 1 0 4h-.1a1.7 1.7 0 0 0-1.5 1z"/></svg>Settings</a>
  </nav>
  <div class="running-pill" id="pill" onclick="go('search')"><b>Search running</b><br><span id="pillTxt"></span></div>
</aside>
<main><div class="wrap">

<section id="p-dash">
  <h1>Dashboard</h1><p class="sub">Your IT company database at a glance.</p>
  <div class="tiles" id="tiles"></div>
  <div class="two">
    <div class="card"><h2>Top cities</h2><div id="topCities"></div></div>
    <div class="card"><h2>Recent searches</h2><div id="recentRuns"></div></div>
  </div>
  <div class="row"><button class="btn" onclick="go('search')">Start a new search</button>
    <a class="btn sec" href="/api/companies.xlsx">Download full database (Excel)</a>
    <a class="btn sec" href="/api/companies.csv">CSV</a>
    <a class="btn sec" href="/api/companies.xlsx?email_type=HR">HR emails only (Excel)</a></div>
</section>

<section id="p-search">
  <h1>New Search</h1><p class="sub">Pick a state and city. The tool collects IT companies from business directories, web search, OpenStreetMap and Google Places (if you add a key), then checks each website's careers, contact and job pages for HR emails.</p>
  <div class="card">
    <div class="grid">
      <div><label>State</label><select id="state"></select></div>
      <div><label>City</label><input id="city" list="cities" placeholder="Type or pick a city"><datalist id="cities"></datalist></div>
      <div><label>Websites to check (max 500)</label><input id="limit" type="number" min="5" max="500"></div>
      <div><label>Min trust score (0-100)</label><input id="trust" type="number" min="0" max="100"></div>
    </div>
    <div class="opts">
      <label class="check"><input type="checkbox" id="skip" checked> Skip companies already in database</label>
      <label class="check"><input type="checkbox" id="hronly"> Only keep companies with an HR email</label>
      <label class="check"><input type="checkbox" id="guess" checked> Suggest hr@/careers@ when none is published</label>
    </div>
    <div class="row mt">
      <button class="btn" id="goBtn">Find companies</button>
      <button class="btn danger" id="stopBtn" style="display:none">Stop</button>
      <span class="right muted" id="srcInfo"></span>
    </div>
  </div>
  <div class="card" id="run" style="display:none">
    <div class="row"><b id="status">Running...</b></div>
    <div class="bar"><i id="bar"></i></div>
    <div class="counts" id="counts"></div>
    <details id="logBox"><summary>Show activity log</summary><div class="log" id="log"></div></details>
  </div>
  <div class="dlbar" id="runDl" style="display:none">
    <b>Download results:</b>
    <a class="btn ok sm" id="dlX">Excel (all)</a>
    <a class="btn sec sm" id="dlXhr">Excel (HR only)</a>
    <a class="btn sec sm" id="dlC">CSV</a>
    <button class="btn sec sm" onclick="copyHr(JOB_ROWS)">Copy all HR emails</button>
  </div>
  <div id="runTable"></div>
</section>

<section id="p-db">
  <h1>Database</h1><p class="sub">All companies you have collected. HR emails are listed first. Filter, copy emails, or export what you see.</p>
  <div class="card">
    <div class="grid">
      <div><label>Search</label><input id="fq" placeholder="Company, website or email"></div>
      <div><label>State</label><select id="fstate"><option value="">All states</option></select></div>
      <div><label>City</label><select id="fcity"><option value="">All cities</option></select></div>
      <div><label>Email type</label><select id="ftype"><option value="">All</option><option>HR</option><option>General</option><option>Guess</option></select></div>
      <div><label>Min trust score</label><input id="ftrust" type="number" min="0" max="100" placeholder="0"></div>
    </div>
    <div class="row mt"><span class="muted" id="dbCount"></span>
      <span class="right"></span>
      <button class="btn sec" onclick="copyDbHr()">Copy HR emails</button>
      <a class="btn sec" id="dbCsv">CSV</a>
      <a class="btn ok" id="dbXlsx">Export to Excel</a></div>
  </div>
  <div id="dbTable"></div>
  <div class="pager" id="pager"></div>
</section>

<section id="p-history">
  <h1>History</h1><p class="sub">Every search you have run.</p>
  <div class="card" id="historyList"></div>
</section>

<section id="p-settings">
  <h1>Settings</h1><p class="sub">Saved on this computer only.</p>
  <div class="card">
    <h2>Google Places API key</h2>
    <p class="muted" style="margin-top:0">Optional. The tool already uses free directories and web search. A Google key adds more companies with phone numbers and addresses.</p>
    <div class="row"><input id="gkey" type="password" placeholder="Paste your key" style="max-width:420px"><button class="btn" onclick="saveKey()">Save key</button><button class="btn danger" onclick="removeKey()">Remove</button></div>
    <p class="hint" id="keyState"></p>
    <p class="hint">Get a key: Google Cloud Console → create project → enable "Places API (New)" → Credentials → Create API key.</p>
  </div>
  <div class="card">
    <h2>Search defaults</h2>
    <div class="grid" style="max-width:420px"><div><label>Websites to check</label><input id="dlimit" type="number" min="5" max="500"></div>
    <div><label>Min trust score</label><input id="dtrust" type="number" min="0" max="100"></div></div>
    <div class="row mt"><button class="btn" onclick="saveDefaults()">Save defaults</button></div>
  </div>
</section>

</div></main>
<div class="toast" id="toast"></div>
<script>
const $=id=>document.getElementById(id);
const esc=s=>String(s??"").replace(/[&<>"']/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]));
const link=(u,t)=>u&&u.startsWith("http")?`<a href="${esc(u)}" target="_blank" rel="noopener">${t}</a>`:'<span class="muted">-</span>';
const host=u=>{try{return new URL(u).hostname.replace(/^www\./,"")}catch(e){return u||""}};
const api=(u,o)=>fetch(u,o).then(r=>r.json());
const post=(u,b)=>api(u,{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(b)});
function toast(t){$("toast").textContent=t;$("toast").classList.add("on");setTimeout(()=>$("toast").classList.remove("on"),1800)}
function copy(t){navigator.clipboard.writeText(t).then(()=>toast(t.length>60?"Copied":"Copied "+t))}
function copyHr(rows){const e=rows.filter(r=>r.email_type==="HR").map(r=>r.email_1);if(!e.length)return toast("No HR emails yet");navigator.clipboard.writeText(e.join(", ")).then(()=>toast(`Copied ${e.length} HR emails`))}
async function copyDbHr(){const d=await api("/api/hr-emails?"+qs());if(!d.emails.length)return toast("No HR emails in this view");navigator.clipboard.writeText(d.emails.join(", ")).then(()=>toast(`Copied ${d.emails.length} HR emails`))}
let STATES={},SETTINGS={},jobId=null,page=1,JOB_ROWS=[];

function go(p){document.querySelectorAll("section").forEach(s=>s.classList.toggle("on",s.id==="p-"+p));
  document.querySelectorAll("nav a").forEach(a=>a.classList.toggle("on",a.dataset.p===p));
  location.hash=p;({dash:loadDash,db:loadDb,history:loadHistory,settings:loadSettings,search:loadSearch}[p]||(()=>{}))()}
document.querySelectorAll("nav a").forEach(a=>a.onclick=()=>go(a.dataset.p));

function scoreCls(s){return s>=70?"s-hi":s>=45?"s-mid":"s-lo"}
const HR_RE=/(^|[._\-+0-9])(hr[a-z]*|careers?|jobs?|recruit[a-z]*|talent[a-z]*|hiring|hire|people|resumes?|cv|joinus|join|placements?|staffing|openings|apply)([._\-+0-9]|$)/i;
function emailCell(r){
  const list=[r.email_1,r.email_2,r.email_3].filter(Boolean);
  const out=list.map((e,i)=>{const hr=(i===0&&r.email_type==="HR")||HR_RE.test(e.split("@")[0]);
    return `<div class="email ${hr?"hr":""}">${hr?'<span class="badge b-hr">HR</span>':""}<b>${esc(e)}</b><button class="copy" onclick="copy('${esc(e)}')">copy</button></div>`}).join("");
  const g=r.guessed_email?r.guessed_email.split(",").map(x=>x.trim()).map(e=>`<div class="email"><span class="badge b-guess">GUESS</span><span class="guess">${esc(e)}</span><button class="copy" onclick="copy('${esc(e)}')">copy</button></div>`).join(""):"";
  const src=r.email_source&&r.email_source.startsWith("http")?`<div class="small muted">found on ${link(r.email_source,esc(new URL(r.email_source).pathname||"/"))}</div>`:(r.email_source?`<div class="small muted">found via ${esc(r.email_source)}</div>`:"");
  return (out+g+src)||'<span class="muted">-</span>'}
function table(rows,{del}={}){
  if(!rows.length)return '<div class="card empty">No companies yet.</div>';
  return `<div class="tbl"><table><thead><tr><th>Company / website</th><th>HR &amp; contact emails</th><th>Type</th><th>Career page</th><th>Contact</th><th>LinkedIn</th><th>Phone</th><th>City</th><th>Trust</th><th>Notes</th>${del?"<th></th>":""}</tr></thead><tbody>`+
  rows.map(r=>`<tr class="${r.email_type==="HR"?"is-hr":""}"><td><b>${esc(r.company)}</b><br>${link(r.website,esc(host(r.website)))}</td>
  <td>${emailCell(r)}</td>
  <td><span class="badge ${r.email_type==="HR"?"b-hr":r.email_type==="Guess"?"b-guess":"b-gen"}">${esc(r.email_type)}</span></td>
  <td>${link(r.career_page,"Open")}</td><td>${link(r.contact_page,"Open")}</td><td>${link(r.linkedin_page,"Open")}</td>
  <td class="small" style="white-space:nowrap">${esc(r.phone)}</td>
  <td>${esc(r.city)}</td><td class="score ${scoreCls(r.trust_score)}">${r.trust_score}</td><td class="note">${esc(r.notes)}</td>
  ${del?`<td><button class="btn danger sm" onclick="delRow('${esc(r.domain)}')">Delete</button></td>`:""}</tr>`).join("")+"</tbody></table></div>"}

async function loadDash(){
  const s=await api("/api/stats");
  $("tiles").innerHTML=[[s.total,"Companies in database"],[s.hr,"With HR email","hr"],[s.career,"With careers page"],[s.cities,"Cities covered"]]
    .map(([n,t,c])=>`<div class="tile ${c||""}"><b>${n}</b><span>${t}</span></div>`).join("");
  $("topCities").innerHTML=s.top_cities.length?`<table><tr><th>City</th><th>Companies</th><th>HR emails</th></tr>${s.top_cities.map(c=>`<tr><td>${esc(c.city)}, <span class="muted">${esc(c.state)}</span></td><td>${c.n}</td><td>${c.hr}</td></tr>`).join("")}</table>`:'<p class="muted">No data yet. Run your first search.</p>';
  $("recentRuns").innerHTML=runsTable(s.runs.slice(0,6));
}
function runsTable(runs){return runs.length?`<table><tr><th>Date</th><th>City</th><th>Checked</th><th>Saved</th><th>HR</th><th>Source</th></tr>${runs.map(r=>`<tr><td>${esc(r.started_at)}</td><td>${esc(r.city)}, <span class="muted">${esc(r.state)}</span></td><td>${r.found}</td><td><b>${r.saved}</b></td><td><b style="color:var(--ok)">${r.hr??"-"}</b></td><td class="muted small">${esc(r.source)}</td></tr>`).join("")}</table>`:'<p class="muted">No searches yet.</p>'}
async function loadHistory(){const s=await api("/api/stats");$("historyList").innerHTML=runsTable(s.runs)}

async function loadSettings(){SETTINGS=await api("/api/settings");
  $("keyState").textContent=SETTINGS.google_key_set?`Key saved (${SETTINGS.google_key_hint}). Searches also use Google Places.`:"No key saved. Searches use directories, web search and OpenStreetMap.";
  $("dlimit").value=SETTINGS.limit;$("dtrust").value=SETTINGS.min_trust;$("gkey").value=""}
async function saveKey(){if(!$("gkey").value.trim())return toast("Paste a key first");await post("/api/settings",{google_key:$("gkey").value});toast("Key saved");loadSettings()}
async function removeKey(){await post("/api/settings",{google_key:""});toast("Key removed");loadSettings()}
async function saveDefaults(){await post("/api/settings",{limit:$("dlimit").value,min_trust:$("dtrust").value});toast("Defaults saved");loadSettings()}

async function loadSearch(){
  SETTINGS=await api("/api/settings");
  if(!$("limit").value)$("limit").value=SETTINGS.limit;if(!$("trust").value)$("trust").value=SETTINGS.min_trust;
  $("srcInfo").innerHTML=SETTINGS.google_key_set?"Sources: Directories + Web + OSM + Google Places":'Sources: Directories + Web search + OSM (<a onclick="go(\'settings\')" style="cursor:pointer">add Google key</a>)';
  if(!Object.keys(STATES).length){STATES=await api("/api/states");$("state").innerHTML=Object.keys(STATES).map(k=>`<option>${k}</option>`).join("");fillCities()}
}
function fillCities(){const c=STATES[$("state").value]||[];$("cities").innerHTML=c.map(x=>`<option value="${x}">`).join("");$("city").value=c[0]||""}
$("state").onchange=fillCities;
$("goBtn").onclick=async()=>{
  if(!$("city").value.trim())return toast("Enter a city");
  const r=await post("/api/start",{state:$("state").value,city:$("city").value,limit:$("limit").value,min_trust:$("trust").value,
    skip_known:$("skip").checked,hr_only:$("hronly").checked,guess:$("guess").checked});
  if(r.error)return toast(r.error);jobId=r.job_id;$("runTable").innerHTML="";poll();
};
$("stopBtn").onclick=async()=>{if(jobId){await post(`/api/job/${jobId}/stop`,{});toast("Stopping... finishing sites already open")}};
async function poll(){
  if(!jobId)return;const j=await api("/api/job/"+jobId);JOB_ROWS=j.rows;
  const running=j.status==="running";
  $("run").style.display="block";$("goBtn").disabled=running;$("stopBtn").style.display=running?"inline-flex":"none";
  const atBottom=$("log").scrollHeight-$("log").scrollTop-$("log").clientHeight<30;
  $("log").textContent=j.log.join("\n");if(atBottom)$("log").scrollTop=1e9;
  $("bar").style.width=(!running?100:j.total?Math.round(j.done*100/j.total):3)+"%";
  $("status").textContent=!running?`Done: ${j.rows.length} companies saved from ${j.city}, ${j.hr} with HR email`:
    (j.discovery_done?`Checking websites in ${j.city}`:`Finding companies in ${j.city} and checking websites`);
  $("counts").innerHTML=[["Websites found",j.total],["Checked",j.done],["Saved",j.rows.length],["With HR email",j.hr,"g"],["Skipped",j.skipped]]
    .map(([t,n,c])=>`<span class="${c||""}">${t}: <b>${n}</b></span>`).join("");
  const cur=$("runTable").dataset.n;if(cur!=String(j.rows.length)){$("runTable").innerHTML=j.rows.length?table(j.rows):"";$("runTable").dataset.n=j.rows.length}
  $("runDl").style.display=j.rows.length?"flex":"none";
  $("dlX").href=`/api/job/${jobId}/export.xlsx`;$("dlXhr").href=`/api/job/${jobId}/export.xlsx?hr_only=1`;$("dlC").href=`/api/job/${jobId}/export.csv`;
  $("pill").style.display=running?"block":"none";$("pillTxt").textContent=`${j.city}: ${j.done}/${j.total||"?"} · ${j.hr} HR`;
  if(running)setTimeout(poll,1200);
}

let FILTERS={};
async function loadDb(){
  FILTERS=await api("/api/filters");const cur=$("fstate").value;
  $("fstate").innerHTML='<option value="">All states</option>'+Object.keys(FILTERS).map(s=>`<option ${s===cur?"selected":""}>${esc(s)}</option>`).join("");
  fillFCities(true);refreshDb();
}
function fillFCities(keep){const cur=keep?$("fcity").value:"";const list=$("fstate").value?FILTERS[$("fstate").value]||[]:[...new Set(Object.values(FILTERS).flat())];
  $("fcity").innerHTML='<option value="">All cities</option>'+list.map(c=>`<option ${c===cur?"selected":""}>${esc(c)}</option>`).join("")}
function qs(){return new URLSearchParams({q:$("fq").value,state:$("fstate").value,city:$("fcity").value,email_type:$("ftype").value,min_trust:$("ftrust").value}).toString()}
async function refreshDb(){
  const d=await api(`/api/companies?${qs()}&page=${page}`);
  $("dbCount").textContent=`${d.total} companies`;$("dbCsv").href="/api/companies.csv?"+qs();$("dbXlsx").href="/api/companies.xlsx?"+qs();
  $("dbTable").innerHTML=table(d.rows,{del:true});
  $("pager").innerHTML=d.pages>1?`<button class="btn sec sm" ${page<=1?"disabled":""} onclick="page--;refreshDb()">Previous</button><span class="muted">Page ${d.page} of ${d.pages}</span><button class="btn sec sm" ${page>=d.pages?"disabled":""} onclick="page++;refreshDb()">Next</button>`:"";
}
let t;$("fq").oninput=()=>{clearTimeout(t);t=setTimeout(()=>{page=1;refreshDb()},300)};
$("fstate").onchange=()=>{fillFCities(false);page=1;refreshDb()};
["fcity","ftype","ftrust"].forEach(id=>$(id).onchange=()=>{page=1;refreshDb()});
async function delRow(d){if(!confirm("Delete "+d+" from the database?"))return;await fetch("/api/companies/"+encodeURIComponent(d),{method:"DELETE"});toast("Deleted");refreshDb()}

(async()=>{const a=await api("/api/job/active");if(a){jobId=a.id;poll()}go(location.hash.slice(1)||"dash")})();
</script></body></html>"""


def free_port(start=5000):
    import socket
    for port in range(start, start + 20):
        with socket.socket() as s:
            if s.connect_ex(("127.0.0.1", port)) != 0:
                return port
    return start


if __name__ == "__main__":
    import webbrowser
    init_db()
    port = free_port()
    url = f"http://127.0.0.1:{port}"
    print(f"\n  HR Finder is running at {url}\n  Keep this window open. Close it to stop.\n")
    threading.Timer(1.5, lambda: webbrowser.open(url)).start()
    app.run(host="127.0.0.1", port=port, debug=False, threaded=True)
